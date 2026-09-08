"""Build editable software-copyright forms from the provided reference forms.

The reference pack contains three legacy .doc forms and one native .docx form.
The legacy forms are first converted to docx outside this script during the
current build and are then used as layout-preserving templates.  This script
keeps the reference table structure and replaces only project-specific content.
Applicant, contact, registration-number and signature fields are intentionally
left for the submitting organisation to complete.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.shared import Pt


ROOT = Path(r"C:\Workspace\Intelligent_Fracturing_Prediction")
QA_TEMPLATES = ROOT / ".docx_qa_20260830" / "example_forms_converted2"
EXAMPLE_DIR = ROOT / "专利与软著" / "exapmle" / "知识图谱软著 (2)"
TOP_OUTPUT = ROOT / "专利与软著" / "deliverables" / "软著"
PACKAGE_OUTPUT = (
    ROOT
    / "专利与软著"
    / "deliverables"
    / "software_copyright_part2_part3"
    / "word_submission"
)

SOFTWARE_NAME = "智能压裂双场景数字孪生与安全建议软件"
SHORT_NAME = "智能压裂双场景数字孪生软件"
VERSION = "V1.0"

FUNCTIONS_TEXT = (
    "软件用途：本软件面向多簇压裂施工数据，用于施工曲线、压力观测、DAS/FracMonitor分簇解释结果、"
    "井段工况和裂缝状态的统一展示与分析，为施工解释、风险识别和参数建议提供数字化工具。\n"
    "\n"
    "软件功能与技术特点：\n"
    "（1）数据与场景管理：加载施工压力、排量、砂比、累计液量、阶段工况、井轨迹及簇几何配置，"
    "并根据观测覆盖范围切换有DAS分簇观测场景和无DAS压力校正场景。\n"
    "（2）压力换算与质量检查：结合静液柱、管柱摩阻、射孔摩阻和压力偏置，将井口压力换算为井底压力，"
    "同时检查缺失值、无效值、时间范围和观测来源。\n"
    "（3）PKN裂缝状态估计：根据总排量、施工阶段和物理参数估计净压力、裂缝半长、缝宽、裂缝体积及六簇分配状态。\n"
    "（4）KG-EnKF参数更新：支持普通EnKF、规则不确定性调节、软先验和相关软先验模式，利用压力及可用分簇观测"
    "更新压力/裂缝参数、簇进液能力和簇间分配参数，并输出先验、后验及更新记录。\n"
    "（5）双场景裂缝展示：有DAS场景展示压力、分簇响应与阶段级裂缝状态；无DAS场景使用压力校正结果驱动"
    "六簇阶段级裂缝半长、缝宽和体积演化，并支持井段切换和时间轴回放。\n"
    "（6）分簇响应与均衡度：计算首次响应、响应效率、时间响应速率、分簇份额、均衡度、Gini系数和归一化熵，"
    "并记录指标的数据来源和状态。\n"
    "（7）阶段总液量建议：通过Piggy-Bank账本记录释放、储备、支取、实际执行量、守恒残差和回滚信息；系统只建议"
    "下一时间窗阶段总排量或总液量的调整，不假设能够直接控制单个簇的实际进液量。\n"
    "（8）风险判断与智能体建议：综合压力安全、异常风险、裂缝改造效果、簇间均衡度和施工成本，输出排量/砂比建议、"
    "风险状态、不确定性和人工确认状态。\n"
    "（9）可视化与导出：提供多井段工况曲线、压力对比、参数更新、六簇响应、三维裂缝演化、风险状态、建议动作、"
    "审计记录和结果导出。\n"
    "\n"
    "运行环境：Windows 10/11，Python 3.11/3.12及项目依赖环境；建议CPU 8核/16线程或以上、内存16GB及以上。\n"
    "编程语言：Python；软件版本号：V1.0；源代码统计：以随附源代码提交件实际统计结果为准。"
)


def all_paragraphs(doc: Document):
    """Yield top-level and table paragraphs, including nested tables."""

    for p in doc.paragraphs:
        yield p

    def walk_table(table):
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
                for nested in cell.tables:
                    yield from walk_table(nested)

    for table in doc.tables:
        yield from walk_table(table)


def set_paragraph_text(paragraph, text: str, font_size: float | None = None):
    """Replace paragraph text while retaining the paragraph style."""

    first_rpr = None
    if paragraph.runs:
        first_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs[0]._r.rPr is not None else None
    paragraph.text = text
    if paragraph.runs:
        run = paragraph.runs[0]
        if first_rpr is not None:
            run._r.insert(0, deepcopy(first_rpr))
        run.font.name = "宋体"
        if font_size is not None:
            run.font.size = Pt(font_size)


def set_cell_text(cell, text: str, font_size: float | None = None):
    if not cell.paragraphs:
        cell.add_paragraph()
    set_paragraph_text(cell.paragraphs[0], text, font_size=font_size)
    for p in cell.paragraphs[1:]:
        set_paragraph_text(p, "", font_size=font_size)


def replace_everywhere(doc: Document, old: str, new: str):
    for p in all_paragraphs(doc):
        if old in p.text:
            set_paragraph_text(p, p.text.replace(old, new))


def set_table_row_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    for child in list(tr_pr):
        if child.tag.endswith("tblHeader"):
            return
    from docx.oxml import OxmlElement

    header = OxmlElement("w:tblHeader")
    header.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
    tr_pr.append(header)


def save_both(doc: Document, top_name: str, package_name: str | None = None):
    TOP_OUTPUT.mkdir(parents=True, exist_ok=True)
    PACKAGE_OUTPUT.mkdir(parents=True, exist_ok=True)
    top_path = TOP_OUTPUT / top_name
    doc.save(top_path)
    package_path = PACKAGE_OUTPUT / (package_name or top_name)
    shutil.copyfile(top_path, package_path)
    return top_path, package_path


def build_declaration():
    doc = Document(QA_TEMPLATES / "附件1-申报书.docx")
    for p in doc.paragraphs:
        if p.text.startswith("计算机软件名称："):
            set_paragraph_text(p, f"计算机软件名称： {SOFTWARE_NAME}")
        elif p.text.startswith("申报单位："):
            set_paragraph_text(p, "申报单位： 待单位确认")
        elif p.text.startswith("申报时间："):
            set_paragraph_text(p, "申报时间： 待填写")
    for old, new in {
        "油气压裂知识图谱交互系统": SOFTWARE_NAME,
        "压裂知识图谱交互系统": SHORT_NAME,
        "中国石化石油工程技术研究院": "待单位确认",
        "中石化石油工程技术研究院有限公司": "待单位确认",
        "中国石油化工股份有限公司": "待确认申报主体",
        "卞晓冰": "待填写",
        "010-56606433": "待填写",
        "8102行": "以提交源代码实际统计为准",
    }.items():
        replace_everywhere(doc, old, new)
    if len(doc.paragraphs) > 2:
        set_paragraph_text(doc.paragraphs[2], "申报单位名称（待填写）")
    # The large application table is the substantive part of the declaration.
    if len(doc.tables) >= 2:
        t0, t1 = doc.tables[0], doc.tables[1]
        set_cell_text(t0.cell(0, 1), SOFTWARE_NAME)
        set_cell_text(t0.cell(1, 1), "待单位确认")
        set_cell_text(t0.cell(2, 1), "待确认著作权人")
        set_cell_text(t0.cell(5, 1), FUNCTIONS_TEXT)
        set_cell_text(t1.cell(0, 2), "待填写")
        set_cell_text(t1.cell(0, 4), "待填写")
    return save_both(doc, "附件1-申报书.docx")


def build_application():
    doc = Document(QA_TEMPLATES / "申请表.docx")
    replace_everywhere(doc, "油气压裂知识图谱交互系统", SOFTWARE_NAME)
    replace_everywhere(doc, "压裂知识图谱交互系统", SHORT_NAME)
    replace_everywhere(doc, "中国石油化工股份有限公司", "待确认著作权人")
    replace_everywhere(doc, "北京市朝阳区朝阳门北大街22号", "待填写")
    replace_everywhere(doc, "01059968825", "待填写")
    if len(doc.tables) >= 2:
        t0, t1 = doc.tables[0], doc.tables[1]
        set_cell_text(t0.cell(0, 3), SOFTWARE_NAME)
        set_cell_text(t0.cell(1, 3), SHORT_NAME)
        set_cell_text(t0.cell(1, 7), f" {VERSION}")
        set_cell_text(t0.cell(5, 2), "待确认著作权人")
        set_cell_text(t0.cell(5, 4), "待确认")
        set_cell_text(t0.cell(5, 6), "待填写")
        set_cell_text(t1.cell(1, 1), "软件用途：\n" + FUNCTIONS_TEXT)
        set_cell_text(t1.cell(2, 1), "软件功能与技术特点：\n" + FUNCTIONS_TEXT)
        set_cell_text(
            t1.cell(3, 1),
            "软件运行软硬件环境：\n软件环境：Windows 10/11，Python 3.11/3.12及项目依赖环境。\n"
            "硬件环境：建议CPU 8核/16线程或以上、内存16GB及以上。",
        )
        set_cell_text(t1.cell(4, 1), "软件编程语言：Python")
        set_cell_text(t1.cell(5, 1), f"软件版本号：{VERSION}")
        set_cell_text(t1.cell(6, 1), "源代码统计：以随附源代码提交件实际统计结果为准")
        set_cell_text(t1.cell(7, 3), "待填写")
        set_cell_text(t1.cell(9, 3), "待填写")
    return save_both(doc, "申请表.docx")


def build_approval():
    doc = Document(QA_TEMPLATES / "审批单.docx")
    for i, p in enumerate(doc.paragraphs):
        if i == 0:
            set_paragraph_text(p, "申报单位名称（待填写）")
        elif i == 2:
            set_paragraph_text(p, "                                                       编号：2026GCY-RZ-________")
    for i, text in {
        6: "软件用途：本软件面向多簇压裂施工数据，为施工解释、风险识别、裂缝状态估计和阶段总液量建议提供数字化工具。",
        9: "本软件将施工曲线、井口—井底压力换算、DAS/FracMonitor分簇解释结果、井段工况和PKN裂缝状态统一组织，形成可回放、可追溯的分析结果。",
        10: "施工数据与场景管理：加载压力、排量、砂比、累计液量、阶段信息、井轨迹及簇几何配置，支持有DAS和无DAS场景切换。",
        11: "压力换算与PKN状态估计：结合静液柱、管柱摩阻、射孔摩阻和压力偏置换算井底压力，并估计净压力、裂缝半长、缝宽和六簇分配。",
        12: "KG-EnKF参数更新：根据压力及可用分簇观测更新压力/裂缝参数、簇进液能力和簇间分配参数，输出先验、后验及更新记录。",
        13: "双场景与分簇可视化：展示压力对比、参数更新、六簇响应、阶段级裂缝演化、均衡度和时间轴回放，并支持结果导出。",
        15: "数据场景分层：按观测覆盖范围区分有DAS分簇观测与无DAS压力校正，缺失观测不以零值替代。",
        16: "物理模型与同化分层：PKN提供快速状态响应，KG-EnKF利用观测更新参数，安全规则对高风险动作进行审核。",
        17: "分簇响应评价：计算首次响应、响应效率、时间响应速率、分簇份额、均衡度、Gini系数和归一化熵。",
        18: "阶段级液量建议：Piggy-Bank只调整下一时间窗阶段总排量或总液量权重，记录释放、储备、支取、守恒、确认和回滚。",
        19: "智能体建议：综合压力安全、异常风险、裂缝几何效果、均衡度和成本输出排量/砂比建议，异常时回退安全动作。",
        20: "结果可追溯：保存输入数据清单、配置、模型模式、参数更新记录、审计记录和导出结果。",
        22: "CPU：8核/16线程或以上；",
        23: "操作系统：Windows 10/11；",
        24: "内存：16GB及以上；",
        25: "硬盘空间：建议至少500GB可用空间；",
        27: "Python",
        29: VERSION,
        31: "以随附源代码提交件实际统计结果为准",
    }.items():
        if i < len(doc.paragraphs):
            set_paragraph_text(doc.paragraphs[i], text)
    replace_everywhere(doc, "页码为：0", "页码为：前、中、后各30页（待核定）")
    if doc.tables:
        t = doc.tables[0]
        set_cell_text(t.cell(0, 1), SOFTWARE_NAME)
        set_cell_text(t.cell(1, 1), "待单位确认")
        set_cell_text(t.cell(2, 1), "待填写")
        set_cell_text(t.cell(2, 3), "待填写")
        set_cell_text(t.cell(3, 1), "待填写")
    return save_both(doc, "审批单.docx")


def build_headquarters_register():
    doc = Document(EXAMPLE_DIR / "登记审批表（总部）.docx")
    replace_everywhere(doc, "压裂知识图谱问答软件", SOFTWARE_NAME)
    replace_everywhere(doc, "油气压裂知识图谱交互系统", SOFTWARE_NAME)
    replace_everywhere(doc, "中国石油化工股份有限公司", "待确认著作权人")
    replace_everywhere(doc, "010-59968836", "待填写")
    replace_everywhere(doc, "章 朋", "待填写")
    if doc.tables:
        t = doc.tables[0]
        set_cell_text(t.cell(0, 2), SOFTWARE_NAME)
        set_cell_text(t.cell(2, 2), "待确认著作权人")
        set_cell_text(t.cell(2, 3), "待确认")
        set_cell_text(t.cell(2, 4), "待填写")
        set_cell_text(t.cell(5, 2), SOFTWARE_NAME)
        set_cell_text(t.cell(6, 1), FUNCTIONS_TEXT)
    if len(doc.paragraphs) > 1:
        set_paragraph_text(doc.paragraphs[1], "编号：                             收案时间：        年     月      日")
    return save_both(doc, "登记审批表（总部）.docx")


def main():
    outputs = [
        build_declaration(),
        build_application(),
        build_approval(),
        build_headquarters_register(),
    ]
    for top, package in outputs:
        print(f"{top}\n{package}")


if __name__ == "__main__":
    main()
