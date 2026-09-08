from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from pathlib import Path


OUT = Path(r"C:\Workspace\Intelligent_Fracturing_Prediction\deliverables\心脏手术后血管麻痹智能辅助识别技术方案_示例.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

FONT_LATIN = "Calibri"
FONT_CJK = "微软雅黑"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
GRAY = "555555"
MUTED = "6B7280"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:{}".format(m)))
        if node is None:
            node = OxmlElement("w:{}".format(m))
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(i, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name=FONT_LATIN, size=11, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size=11, color=BLACK, bold=False, italic=False):
    style.font.name = FONT_LATIN
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic


def set_paragraph_spacing(p, before=0, after=8, line=1.333, align=None, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True


def add_num_definition(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(existing_abs or [0]) + 1
    num_id = max(existing_num or [0]) + 1
    abs_num = OxmlElement("w:abstractNum")
    abs_num.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abs_num.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet")
    lvl.append(fmt)
    txt = OxmlElement("w:lvlText")
    txt.set(qn("w:val"), "•")
    lvl.append(txt)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    lvl.append(ppr)
    abs_num.append(lvl)
    numbering.append(abs_num)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abs_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(doc._bullet_num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p._p.get_or_add_pPr().append(num_pr)
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_para(doc, text="", style="Normal", bold_prefix=None, after=None, before=None, align=None, color=None, size=None):
    p = doc.add_paragraph(style=style)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=size or 11, color=color or BLACK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=size or 11, color=color or BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size or 11, color=color or BLACK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, size={1:16,2:13,3:12}[level], color={1:BLUE,2:BLUE,3:DARK_BLUE}[level], bold=True)
    return p


def add_callout(doc, label, text, fill=CALLOUT, label_color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val":"single","sz":"6","color":"D7DBE2"}, bottom={"val":"single","sz":"6","color":"D7DBE2"}, left={"val":"single","sz":"14","color":BLUE}, right={"val":"single","sz":"6","color":"D7DBE2"})
    p = cell.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(label + "：")
    set_run_font(r, size=10.5, color=label_color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def format_table_text(cell, text, header=False, color=BLACK, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if not header else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    set_run_font(r, size=size, color=color, bold=header)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        format_table_text(hdr.cells[i], h, header=True, color=NAVY, size=font_size)
        set_cell_shading(hdr.cells[i], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            format_table_text(cells[i], value, header=False, size=font_size)
            set_cell_shading(cells[i], WHITE)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={"val":"single","sz":"4","color":"D7DBE2"}, bottom={"val":"single","sz":"4","color":"D7DBE2"}, left={"val":"single","sz":"4","color":"D7DBE2"}, right={"val":"single","sz":"4","color":"D7DBE2"})
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def add_metadata_table(doc):
    rows = [
        ["研究类型", "单中心回顾性观察性研究；预测模型开发与内部验证", "数据规模", "约400例，最终以清洗后有效病例数为准"],
        ["研究对象", "成人体外循环心脏手术患者（纳排标准待临床团队确认）", "模型时点", "术后进入ICU即刻"],
        ["主要任务", "识别术后早期血管麻痹表型/风险", "当前用途", "研究性辅助识别，不作独立临床诊断"],
        ["版本属性", "技术方案示例；待数据核验、标签复核和伦理审批", "编制日期", "2026年8月"],
    ]
    return add_table(doc, ["项目项", "方案设定", "项目项", "方案设定"], rows, [1500, 3200, 1500, 3160], header_fill=LIGHT_GRAY, font_size=9.3)


def build_document():
    doc = Document()
    doc._bullet_num_id = add_num_definition(doc)
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 11, BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level, size, color, before, after in [(1,16,BLUE,18,10),(2,13,BLUE,12,6),(3,12,DARK_BLUE,8,4)]:
        st = styles[f"Heading {level}"]
        set_style_font(st, size, color, True)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.keep_with_next = True
    title = styles["Title"]
    set_style_font(title, 24, NAVY, True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = styles["Subtitle"]
    set_style_font(subtitle, 13, GRAY, False)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(20)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("心脏手术后血管麻痹 | 技术方案示例")
    set_run_font(r, size=8.5, color=MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    r = footer.add_run("研究方案示例  |  第 ")
    set_run_font(r, size=9, color=MUTED)
    add_page_field(footer)
    r = footer.add_run(" 页")
    set_run_font(r, size=9, color=MUTED)

    # Opening block
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("研究技术方案")
    set_run_font(r, size=11, color=BLUE, bold=True)
    p = doc.add_paragraph(style="Title")
    r = p.add_run("心脏手术后血管麻痹智能辅助识别")
    set_run_font(r, size=24, color=NAVY, bold=True)
    p = doc.add_paragraph(style="Subtitle")
    r = p.add_run("基于约400例围术期结构化数据的示例方案")
    set_run_font(r, size=13, color=GRAY)
    add_metadata_table(doc)
    add_callout(doc, "先行结论", "现有数据适合先完成“研究性辅助识别模型”的开发、内部验证和变量体系标准化；在补齐连续MAP/血压、实际血管活性药物剂量与时间、CO/CI或SVR、尿量和结局信息前，不建议表述为实时预警系统或独立诊断工具。")

    add_heading(doc, "一、项目定位与研究目标", 1)
    add_heading(doc, "1.1 项目定位", 2)
    add_para(doc, "本项目拟建立一个面向成人体外循环心脏手术患者的围术期血管麻痹智能辅助识别流程。模型的输入来自术前基础状态、术中手术暴露以及术后进入ICU即刻可获得的检验与生理指标，输出为血管麻痹的预测概率、分层结果和可解释的主要影响因素。")
    add_para(doc, "一期不追求复杂深度学习，而是优先完成可复现、可解释、可审计的研究管线：明确研究时点和标签、清理变量、建立基线模型、进行严格的内部验证，并形成后续扩大样本和外部验证的标准。")
    add_heading(doc, "1.2 研究目标", 2)
    for t in [
        "建立一套适用于本中心数据的血管麻痹表型与标签核查规则，明确标签来源、发生时间窗及不确定病例处理方式。",
        "形成围术期标准化变量字典，完成单位、编码、缺失值、重复字段、异常值和时间顺序的统一。",
        "在术后进入ICU即刻时间点，开发血管麻痹辅助识别模型，并与单一NEE旗标、传统统计模型和简单临床规则进行比较。",
        "用区分度、校准度、临床净获益、稳定性及可解释性评价模型，明确当前模型可以支持的研究用途和不能支持的临床用途。",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "二、研究问题与可行性边界", 1)
    add_heading(doc, "2.1 建议采用的研究问题", 2)
    add_para(doc, "推荐将一期问题写成：在成人体外循环心脏手术患者中，能否利用术前、术中及术后进入ICU即刻的结构化临床数据，辅助识别术后预设时间窗内的血管麻痹表型？")
    add_para(doc, "这里的“识别”是研究性算法输出，不等同于临床诊断。若希望研究“预测”，必须把预测时点、预测窗口和所有输入变量的可获得时间写入方案；若血管麻痹标签依赖术后持续升压药剂量，则该剂量只能用于标签形成或后续动态预警，不能同时作为同一时点的输入特征。")
    add_heading(doc, "2.2 现有数据可做与暂不能做", 2)
    add_table(doc, ["范围", "当前判断", "技术处理"], [
        ["可以完成", "利用现有二分类字段“是否血管麻痹”进行初步模型开发和内部验证。", "先做标签来源审计，再固定分析数据集和模型时点。"],
        ["可以完成", "比较围术期基础病、心功能、手术暴露、乳酸、血常规、肝肾功能、凝血和ICU即刻指标的联合信息。", "按临床域分组，避免把全部原始字段无约束地喂给模型。"],
        ["暂不宜宣称", "实时连续监测、动态报警、独立诊断或软件医疗器械性能。", "补齐时间序列和外部验证后再进入下一阶段。"],
        ["关键缺口", "字段清单中未见实际MAP/血压时间序列、各血管活性药物剂量及持续时间、CO/CI、SVR、尿量和AKI/ICU结局。", "列为数据补采清单，并在论文中如实报告。"],
    ], [1350, 4050, 3960])

    add_heading(doc, "三、研究对象、索引时点与结局标签", 1)
    add_heading(doc, "3.1 研究对象", 2)
    add_para(doc, "拟纳入本中心成人体外循环心脏手术患者。纳入年龄、手术日期范围、首次手术/重复手术、急诊状态、机械循环支持和感染性心内膜炎病例等，应在数据冻结前由临床团队形成书面纳排标准，并对排除病例保留原因编码。")
    add_heading(doc, "3.2 索引时点与预测窗口", 2)
    add_table(doc, ["要素", "建议设定", "必须确认的内容"], [
        ["索引时点", "术后进入ICU即刻；使用ICU即刻检验与当时可获得的围术期信息。", "是否能获得准确的ICU入室日期和时间；“立即”允许的时间范围。"],
        ["结局窗口", "建议示例为进入ICU后24小时内；如病例记录以术后48小时为主，则统一改为48小时。", "“是否血管麻痹”对应的观察窗口和结束时间。"],
        ["输出形式", "连续风险概率 + 预先定义的低/中/高风险层级。", "阈值由验证集/交叉验证和临床使用目的共同确定，不从训练集单独挑选。"],
        ["发生时机", "作为结局的时间分层或敏感性分析变量。", "不能作为ICU即刻模型的输入，除非它在索引时点之前已知。"],
    ], [1500, 3960, 3900])
    add_heading(doc, "3.3 标签体系与标签泄漏规则", 2)
    add_para(doc, "建议建立三级标签，而不是直接把表格中的“是否血管麻痹”当作无条件金标准。所有标签在数据冻结前完成核查，并保留原始标签、复核标签和最终分析标签。")
    add_table(doc, ["标签", "定义/用途", "当前可行性"], [
        ["主分析标签 Y1", "临床团队复核后的“是否血管麻痹”二分类标签；需记录依据、观察窗、是否经过专家判定。", "可利用现有字段，但需追溯来源。"],
        ["操作性标签 Y2", "拟结合NEE≥0.2、低血压/MAP、心功能/灌注状态及出血、填塞、严重低心排等排除因素形成。", "当前缺少实际MAP、剂量和排除信息，不能仅凭NEE旗标独立重建。"],
        ["不确定标签 YU", "信息冲突、观察窗不完整、标签来源不清或关键字段缺失的病例。", "主分析排除或单列；敏感性分析纳入不同处理方案。"],
    ], [1500, 5000, 2860])
    add_callout(doc, "核心防泄漏规则", "如果“是否血管麻痹”本身是由“是否达到NEE≥0.2”计算得到，NEE旗标不能作为该模型的预测变量；如果标签来自独立临床判定，则仍需做“剔除NEE旗标”的敏感性分析。")

    add_heading(doc, "四、现有字段体系与数据字典方案", 1)
    add_heading(doc, "4.1 变量域划分", 2)
    add_table(doc, ["变量域", "拟纳入字段", "模型角色与注意事项"], [
        ["身份与管理字段", "住院号、姓名", "仅用于去重、病例追溯和权限管理；不进入模型，建模前删除或脱敏。"],
        ["人口学与体格", "性别、年龄（y）、身高（cm）、体重（kg）、BMI", "年龄、BMI优先保留连续形式；BMI与身高/体重存在派生关系，避免无目的重复输入。"],
        ["基础疾病", "高血压、糖尿病、高脂血症、吸烟史、饮酒史、慢性肾病、慢性肺病、慢性肝病、甲状腺疾病、脑梗病史、冠心病史、心梗史、心衰史、房颤史、心脏手术史、本次感染性心内膜炎、自身/风湿免疫病史", "统一为0/1/未知；“未知”不能直接当作“无”。感染性心内膜炎和既往心衰等可作为临床重要亚组。"],
        ["心功能与术前用药", "LVEF（%）、左心室前后径（mm）、NYHA分级、ACEI/ARB/ARNI、β受体阻滞剂、钙拮抗药、利尿药、硝酸酯药、正性肌力药", "NYHA为有序分类；药物字段需明确是术前长期使用、入院使用还是手术当天使用。"],
        ["术前实验室", "WBC、中性粒、淋巴粒、NLR、SII、PWR、Hb、PLT、ALT、AST、白蛋白、总胆红素、Cr、eGFR、INR、纤维蛋白原、促甲状腺素、游离T3、游离T4、NT-ProBNP、肌钙蛋白I", "核对单位、检测时间和参考范围；NLR/SII/PWR优先用原始细胞计数复算并与原字段比对。"],
        ["诊断与手术暴露", "诊断、手术类型、手术结束日期、手术时长（min）、CPB时长（min）、术中乳酸最低值、术中乳酸峰值、全血（ml）、术中RBC（U）、术中血小板、凝血因子（U）、术中FFP（ml）", "手术类型在清单中出现两次，需核验是否为重复列或两个不同字段；血液制品统一单位和统计窗口。"],
        ["ICU即刻指标", "术后进入ICU立即pH、钾离子、钙离子、乳酸、Hb、WBC、中性粒、淋巴粒、NLR、SII、PWR、Hb（g/L）、PLT、ALT、AST、白蛋白、总胆红素、Cr、eGFR、INR、纤维蛋白原", "“立即Hb”与“立即Hb（g/L）”疑似重复或单位缺失，需回到原始系统核对。只有在索引时点可获得的值才可进入即刻模型。"],
        ["结局与辅助字段", "是否达到NEE≥0.2、是否血管麻痹、发生时机", "是否血管麻痹为主结局候选；NEE旗标优先作为标签依据/审计变量；发生时机用于结局分层。"],
    ], [1450, 4550, 3360], font_size=8.8)
    add_heading(doc, "4.2 建议新增的最低限度字段", 2)
    add_para(doc, "如果可以从麻醉、ICU和电子病历系统回补数据，建议优先新增以下字段；它们对标签可信度和后续临床转化的价值高于继续增加普通静态变量。")
    for t in [
        "精确时间：手术开始/结束、CPB开始/结束、主动脉阻断、ICU入室、首次低血压、首次达到升压药阈值、血管麻痹判定时间。",
        "血流动力学：MAP/收缩压/舒张压时间序列，至少记录低血压持续时间和索引时点值；如有则记录CO、CI、SVR及监测方式。",
        "血管活性药物：去甲肾上腺素、肾上腺素、血管加压素、多巴胺、去氧肾上腺素等的实际剂量、体重归一化方式、起止时间和累计暴露。",
        "鉴别诊断与结局：出血量/再开胸、容量复苏、低心排、心包填塞、机械循环支持、尿量、AKI、机械通气时间、ICU住院时间和院内死亡。",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "五、数据治理与预处理流程", 1)
    add_heading(doc, "5.1 数据冻结与质量控制", 2)
    add_table(doc, ["步骤", "实施要求", "输出"], [
        ["病例级去重", "以脱敏病例键、手术日期和手术类型联合检查重复；同一患者多次手术明确分析单位。", "病例流图、去重日志、最终样本数。"],
        ["字段审计", "核验列名、中文单位、编码、日期格式、重复列、全空列和无法解释的缩写。", "数据字典v1、字段问题清单。"],
        ["数值质控", "检查负值、极端值、单位错位和不可能组合；异常值先标记，不直接删除。", "异常值报告、人工复核记录。"],
        ["缺失审计", "报告每列缺失率、每病例缺失数、按结局分层的缺失模式，区分未测量与录入缺失。", "缺失热图/表、插补策略。"],
        ["标签审计", "抽取标签阳性、阴性、冲突和边界病例，由至少两名临床人员复核，分歧由第三人裁决。", "标签复核表、最终标签版本。"],
        ["数据冻结", "锁定分析数据集、代码版本、随机种子、变量清单和方案版本后再建模。", "可复现分析包。"],
    ], [1450, 5200, 2710])
    add_heading(doc, "5.2 变量处理原则", 2)
    for t in [
        "连续变量：保留临床合理的连续信息；对NT-ProBNP、肌钙蛋白、胆红素等偏态指标可在训练折内做log1p或稳健变换，并记录变换规则。",
        "分类变量：二分类统一为0/1/未知；手术类型、诊断等高基数变量合并稀有类别，合并规则只能在训练数据中确定。",
        "缺失值：主分析使用训练折内中位数/众数或多重插补；同时加入缺失指示变量。插补器、标准化器和特征筛选器不得在全数据上预先拟合。",
        "派生指标：复算BMI、NLR、SII、PWR并与原列进行一致性核验；若公式或单位不清，保留原列但标记为不可作为主模型特征。",
        "时间约束：任何发生在索引时点之后的数据不能进入“ICU即刻”模型，包括后续乳酸、后续升压药剂量、发生时机和后续治疗反应。",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "六、模型开发技术路线", 1)
    add_heading(doc, "6.1 模型分层", 2)
    add_table(doc, ["模型", "方法", "用途"], [
        ["M0 基线", "结局发生率、单变量规则或单一NEE旗标（仅用于对照，不作为最终模型）。", "判断联合模型是否真正增加识别信息。"],
        ["M1 主模型", "带弹性网正则化的二分类逻辑回归；完整预处理封装在交叉验证管线内。", "首选可解释、可校准、适合约400例小样本的模型。"],
        ["M2 探索模型", "限制深度和叶节点数的梯度提升模型，如CatBoost/LightGBM/XGBoost之一。", "检验非线性与交互是否带来稳定增益；不以单次AUC最高为唯一选择标准。"],
        ["M3 简化模型", "根据临床域和稳定性筛选8–12个左右核心变量的简化逻辑模型。", "形成可解释的研究评分或后续原型接口。"],
    ], [1500, 4900, 2960])
    add_heading(doc, "6.2 特征选择与样本量约束", 2)
    add_para(doc, "约400例数据不适合直接训练深度神经网络，也不适合把全部字段未经约束地用于复杂模型。应先统计血管麻痹事件数E，再根据事件数、缺失率和变量相关性限制模型复杂度。若E较少，优先使用M1和M3，并将结果标记为探索性；只有在事件数和有效样本量支持时，才报告M2的增益。")
    add_callout(doc, "建议的复杂度底线", "最终模型的有效自由度应由事件数和内部验证结果共同决定。可将每10个结局事件支持约1个有效参数作为保守起点，而不是把它当成硬性统计定律；事件数不足时应减少变量、合并类别并扩大样本，而不是依靠更复杂算法弥补。", fill="FFF8E8", label_color="7A5A00")
    add_heading(doc, "6.3 推荐分析管线", 2)
    for t in [
        "固定研究问题、索引时点、结局窗口和变量白名单。",
        "按患者而不是按行拆分数据；若同一患者有多次手术，必须保证同一患者不同时出现在训练折和验证折。",
        "采用重复分层5折交叉验证，建议5次重复；每个训练折内完成缺失处理、变换、类别合并、正则化参数选择和模型拟合。",
        "输出每个病例的折外预测概率，汇总ROC-AUC、PR-AUC、Brier分数和校准指标。",
        "以临床可解释性、校准和稳定性为主选择模型；用bootstrap评估系数、预测概率和特征贡献的稳定性。",
        "若未来病例数足够，优先按时间建立后续病例验证集；当前不能把随机交叉验证称为外部验证。",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "七、性能评价与统计分析", 1)
    add_heading(doc, "7.1 主要评价指标", 2)
    add_table(doc, ["维度", "指标", "报告方式"], [
        ["判别能力", "ROC-AUC、PR-AUC", "报告点估计和95%置信区间；重点关注PR-AUC在低发生率场景下的意义。"],
        ["分类性能", "敏感度、特异度、PPV、NPV、F1或平衡准确率", "阈值预先设定或通过交叉验证确定；同时报告阈值对应的病例数。"],
        ["校准能力", "Brier分数、校准截距、校准斜率、校准曲线", "不能只报告AUC；校准曲线使用折外预测概率。"],
        ["临床应用价值", "决策曲线分析、不同阈值下净获益", "明确阈值对应的临床动作，例如复核或加强监测，而不是直接改变治疗。"],
        ["稳定性", "变量入选频率、系数方向、bootstrap预测波动", "识别样本量不足或标签不稳定造成的虚假重要特征。"],
    ], [1450, 3300, 4610])
    add_heading(doc, "7.2 对照、亚组与敏感性分析", 2)
    for t in [
        "与单一NEE阈值、单一MAP阈值（若补齐MAP）、仅术前变量模型和仅术后即刻变量模型比较。",
        "比较Y1临床标签与Y2操作性标签；若两者不一致，报告一致性和对模型性能的影响。",
        "分别采用完整病例、训练折内单次插补和多重插补进行敏感性分析。",
        "按手术类型、感染性心内膜炎、LVEF分层、肾功能状态和年龄段开展探索性亚组分析；事件数不足的亚组不作确定性结论。",
        "进行标签置换/负对照检查或时间泄漏审计，确认模型没有通过日期、标签派生字段或后验治疗信息间接识别结局。",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "八、可解释性与临床呈现", 1)
    add_heading(doc, "8.1 解释输出", 2)
    add_table(doc, ["输出层", "建议内容"], [
        ["群体层面", "报告模型系数/优势比、变量入选频率、Permutation importance或SHAP汇总图，并区分关联解释与因果解释。"],
        ["个体层面", "输出风险概率、风险层级、前3–5个主要贡献因素、缺失字段提醒和“需临床复核”提示。"],
        ["数据层面", "显示输入数据时间、单位、异常值和缺失情况；不允许模型在字段不完整时静默输出高置信度结果。"],
        ["临床边界", "明确该输出只用于辅助筛查/研究分层，不替代血流动力学判断、床旁超声、容量评估和临床决策。"],
    ], [1900, 7460])
    add_heading(doc, "8.2 原型系统形态", 2)
    add_para(doc, "一期建议交付“离线分析原型”：导入脱敏CSV/Excel后完成字段校验、预处理、风险计算和报告导出。待获得连续数据和前瞻性流程后，再考虑接入数据平台进行定时批量计算或实时接口。系统应保留模型版本、输入快照、输出时间和人工复核记录，便于追溯。")

    add_heading(doc, "九、实施计划与交付物", 1)
    add_table(doc, ["阶段", "主要工作", "交付物"], [
        ["第1阶段：数据审计", "病例数、事件率、重复列、单位、缺失、标签来源和索引时间核验。", "数据审计报告、变量字典v1、问题字段清单。"],
        ["第2阶段：标签复核", "临床双人复核Y1；制定Y2候选规则；处理不确定病例。", "标签判定手册、复核表、最终标签版本。"],
        ["第3阶段：基线分析", "描述性统计、阳性/阴性组比较、缺失模式和单变量基线。", "Table 1、缺失图、基线模型报告。"],
        ["第4阶段：模型开发", "M1主模型、M2探索模型、M3简化模型；所有步骤可复现。", "训练代码、模型文件、折外预测结果。"],
        ["第5阶段：内部验证", "重复分层交叉验证、bootstrap、校准、决策曲线、敏感性和亚组分析。", "性能报告、解释性图表、局限性清单。"],
        ["第6阶段：方案固化", "形成论文/课题报告、数据字典、版本记录和后续补采计划。", "技术方案v2、分析报告、原型接口需求。"],
    ], [1600, 4750, 3010])
    add_heading(doc, "十、风险、伦理与质量边界", 1)
    for t in [
        "样本量和事件数可能不足：当前结果应定位为模型开发与内部验证，不报告未经验证的临床效用，不把一次交叉验证的高AUC作为推广依据。",
        "标签存在定义异质性：需要记录标签来源和观察窗；对于无法确认的病例，宁可保留不确定状态，也不要强行归为阴性。",
        "严重数据泄漏风险：NEE旗标、发生时机、后验治疗、后续实验室和结局变量必须按时间审计。",
        "变量测量偏倚：不同手术类型、监测设备、用药习惯和记录完整度可能造成模型偏倚，应报告亚组缺失和性能差异。",
        "隐私与权限：住院号和姓名只保留在受控的病例映射表中；分析库使用脱敏ID，导出图表不包含直接身份信息。",
        "临床责任边界：模型输出为辅助信息，不能单独触发升压药、容量治疗或其他治疗动作；正式部署前需经过前瞻性验证、临床工作流评估及相应审批。",
    ]:
        add_bullet(doc, t)
    add_callout(doc, "建议的项目名称", "现阶段可使用“基于围术期结构化数据的心脏手术后血管麻痹智能辅助识别模型开发与内部验证”。待补齐时间序列、完成外部验证并开展临床影响评价后，再考虑使用“实时预警”或“辅助诊断工具”等更强表述。")

    add_heading(doc, "十一、结论", 1)
    add_para(doc, "基于目前约400余例和所列字段，最稳妥的技术路径是：先把研究对象、索引时点和血管麻痹标签定义清楚，完成字段质量审计与临床复核；再以弹性网逻辑回归为主、梯度提升为辅，采用重复分层交叉验证和bootstrap开展内部验证，重点报告校准、稳定性和可解释性。该路径能够形成可用于课题申报、伦理讨论和后续扩样的数据基础，同时清楚标示当前数据距离实时临床系统之间的差距。")

    add_heading(doc, "附录A：建模前必须确认的字段问题", 1)
    add_table(doc, ["问题字段/问题", "需要确认的内容", "默认处理建议"], [
        ["手术类型重复出现两次", "是否为同一字段重复导出，还是分别代表术式一级/二级分类。", "未确认前不同时纳入；重命名为手术类型_1/2并保留来源。"],
        ["术后立即Hb出现无单位和g/L两列", "是否同一检测值、是否一列为其他时间点或录入错误。", "核对原始记录；只保留单位明确且时间明确的一列。"],
        ["PWR定义不清", "分子、分母、单位和计算时间点。", "优先由原始细胞计数复算；公式不明时不进入主模型。"],
        ["术中血小板无单位", "是U、治疗次数、袋数还是ml。", "统一为原始单位并在字典中写明。"],
        ["全血/FFP/凝血因子/血小板统计窗口", "术中、手术结束前还是整个住院阶段。", "限定为索引时点前可知的围术期窗口。"],
        ["NEE≥0.2仅为二分类旗标", "是否有实际药物剂量、体重归一化和持续时间。", "作为标签审计字段；补齐实际剂量后再做动态模型。"],
        ["是否血管麻痹来源", "临床判定、病历诊断、科研定义还是由NEE计算。", "必须写入标签来源和判定时间，并保留复核版本。"],
    ], [2250, 4300, 2810])

    add_heading(doc, "附录B：建议的分析数据集字段分层", 1)
    add_table(doc, ["层级", "可用字段示例", "建模规则"], [
        ["X0：身份/管理", "脱敏病例ID、手术日期、数据版本", "仅用于管理、去重和分层，不进入特征矩阵。"],
        ["X1：术前", "人口学、基础病、心功能、术前用药、术前实验室", "用于基线模型和术前风险分层。"],
        ["X2：术中", "术式、CPB、手术时长、乳酸、输血", "仅保留手术结束前已知信息。"],
        ["X3：ICU即刻", "pH、电解质、乳酸、血常规、肝肾功能、凝血", "只有进入ICU即刻窗口内的首次合格测量进入即刻模型。"],
        ["Y：结局", "是否血管麻痹、发生时机、复核标签", "不进入特征矩阵；用于主分析、分层和敏感性分析。"],
        ["E：后验/扩展", "持续MAP、实际升压药剂量、尿量、AKI、ICU结局", "一期缺口；补齐后用于动态模型、结局模型或外部验证。"],
    ], [1500, 4500, 3360])

    add_heading(doc, "附录C：方法学参考", 1)
    add_para(doc, "1. Collins GS, et al. TRIPOD+AI statement: updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ. 2024;385:e078378. https://www.bmj.com/content/385/bmj.q902")
    add_para(doc, "2. Moons KG, et al. PROBAST+AI: an updated quality, risk of bias, and applicability assessment tool for prediction models using regression or artificial intelligence methods. BMJ. 2025;388:e082505. https://www.bmj.com/content/388/bmj-2024-082505")
    add_para(doc, "3. Independent factors for the development of vasoplegic syndrome in patients undergoing coronary artery bypass surgery. 该研究示例将MAP、NEE和中心静脉血氧饱和度等条件组合用于操作性定义，可作为标签设计的参考，但不能替代本中心临床团队对观察窗和排除条件的确认。https://pmc.ncbi.nlm.nih.gov/articles/PMC11420007/")

    # Set document core properties
    props = doc.core_properties
    props.title = "心脏手术后血管麻痹智能辅助识别技术方案（示例）"
    props.subject = "围术期真实世界数据与预测模型开发"
    props.author = ""
    props.comments = "技术方案示例；需结合实际数据和临床标准进一步确认。"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
