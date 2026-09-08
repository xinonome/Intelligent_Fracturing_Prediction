from __future__ import annotations

import math
import hashlib
import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
# The IP package is intentionally kept under the Chinese-named project
# folder.  The repository root also contains unrelated delivery artifacts,
# so the builder must not silently write into that other directory.
DELIVERABLES = ROOT / "专利与软著" / "deliverables"
PATENT = DELIVERABLES / "patent_part2_part3"
SOFTWARE = DELIVERABLES / "software_copyright_part2_part3"
WORD_PATENT = PATENT / "word_submission"
WORD_SOFTWARE = SOFTWARE / "word_submission"
WORD_SOURCE = SOFTWARE / "source_code_submission"
FIGURES = PATENT / "figures" / "rendered"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "666666"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


def set_run_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)


def configure_document(doc, label, preset="standard_business_brief"):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10 if preset == "standard_business_brief" else 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, indent, hanging in [("List Bullet", 720, 360), ("List Number", 720, 360)]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Pt(indent / 20)
        style.paragraph_format.first_line_indent = Pt(-hanging / 20)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    code.font.size = Pt(8.5)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.text = "智能压裂项目 · " + label
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Internal submission draft · Page ")
    add_page_field(footer)
    for run in footer.runs:
        set_run_font(run, size=9, color=MUTED)


def add_title_block(doc, title, subtitle, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(label)
    set_run_font(run, size=11, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(title)
    set_run_font(run, size=24, color=INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(subtitle)
    set_run_font(run, size=12.5, color=MUTED)
    add_metadata_table(doc, [
        ("项目", "智能压裂预测与联合动态联调平台"),
        ("材料版本", "V1.0 · 2026-08-30"),
        ("材料性质", "技术底稿转正式整理版，待主体信息和代理师审查"),
        ("数据边界", "仅含脱敏代码、合成样例和可公开技术描述"),
    ])


def add_metadata_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], LIGHT_BLUE)
        for idx, cell in enumerate(cells):
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                for run in para.runs:
                    set_run_font(run, size=9.5, color=INK if idx == 1 else DARK_BLUE, bold=idx == 0)
    set_table_geometry(table, [1800, 7560])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(doc, text, label="说明"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{label}：")
    set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_md_lines(md_path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    items = []
    in_code = False
    code_lang = ""
    buffer = []
    table_buffer = []
    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            items.append(("table", table_buffer))
            table_buffer = []
    def flush_code():
        nonlocal buffer, in_code, code_lang
        if buffer:
            items.append(("code", (code_lang, buffer)))
            buffer = []
        in_code = False
        code_lang = ""
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_table()
            if in_code:
                flush_code()
            else:
                in_code = True
                code_lang = line[3:].strip()
            continue
        if in_code:
            buffer.append(line)
            continue
        if line.strip().startswith("|"):
            table_buffer.append(line.strip())
            continue
        flush_table()
        if not line.strip():
            items.append(("blank", ""))
        elif re.match(r"^#{1,3}\s+", line):
            m = re.match(r"^(#{1,3})\s+(.*)$", line)
            items.append(("heading", (len(m.group(1)), m.group(2).strip())))
        elif line.startswith("> "):
            items.append(("quote", line[2:].strip()))
        elif re.match(r"^[-*]\s+", line):
            items.append(("bullet", re.sub(r"^[-*]\s+", "", line)))
        elif re.match(r"^\d+[.)]\s+", line):
            items.append(("number", re.sub(r"^\d+[.)]\s+", "", line)))
        elif re.match(r"^[-*_]{3,}$", line.strip()):
            items.append(("rule", ""))
        else:
            items.append(("para", line.strip()))
    flush_table()
    if in_code:
        flush_code()
    return items


def add_table_from_markdown(doc, rows):
    parsed = []
    for row in rows:
        parts = [p.strip() for p in row.strip("|").split("|")]
        parsed.append(parts)
    if len(parsed) > 1 and all(re.match(r"^:?-{3,}:?$", x.replace(" ", "")) for x in parsed[1]):
        parsed.pop(1)
    cols = max(len(r) for r in parsed)
    widths = [9360 // cols] * cols
    widths[-1] += 9360 - sum(widths)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for ridx, row in enumerate(parsed):
        cells = table.add_row().cells
        for idx in range(cols):
            text = row[idx] if idx < len(row) else ""
            cells[idx].text = text
            if ridx == 0:
                set_cell_shading(cells[idx], LIGHT_BLUE)
            for para in cells[idx].paragraphs:
                para.paragraph_format.space_after = Pt(1)
                for run in para.runs:
                    set_run_font(run, size=9, color=DARK_BLUE if ridx == 0 else INK, bold=ridx == 0)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_markdown_body(doc, md_path):
    first_h1_seen = False
    for kind, value in parse_md_lines(md_path):
        if kind == "heading":
            level, text = value
            if level == 1 and not first_h1_seen:
                first_h1_seen = True
                continue
            # The source Markdown title is already represented by the formal
            # title block, so the remaining hierarchy starts at Word Heading 1.
            doc.add_heading(text, level=max(1, level - 1))
        elif kind == "para":
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = False
            r = p.add_run(value)
            set_run_font(r, size=11, color=INK)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(value)
            set_run_font(r, size=11, color=INK)
        elif kind == "number":
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(value)
            set_run_font(r, size=11, color=INK)
        elif kind == "quote":
            add_callout(doc, value, "技术口径")
        elif kind == "code":
            lang, code_lines = value
            if lang and lang not in ("text", "plain"):
                p = doc.add_paragraph()
                r = p.add_run(f"[{lang}]")
                set_run_font(r, size=8.5, color=MUTED, italic=True)
            for code_line in code_lines:
                p = doc.add_paragraph(style="Code Block")
                r = p.add_run(code_line if code_line else " ")
                set_run_font(r, name="Consolas", size=8.5, color=INK)
        elif kind == "rule":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "D7DBE2")
            pBdr.append(bottom)
            pPr.append(pBdr)
        elif kind == "table":
            add_table_from_markdown(doc, value)
        elif kind == "blank":
            continue


def add_toc_note(doc, entries):
    doc.add_heading("材料目录", level=1)
    for label, source in entries:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"{label}（源文件：{source}）")
        set_run_font(r, size=10.5, color=INK)


def build_markdown_doc(source, output, title, subtitle, label, preset="standard_business_brief", intro=None):
    doc = Document()
    configure_document(doc, label, preset=preset)
    add_title_block(doc, title, subtitle, label)
    if intro:
        add_callout(doc, intro, "递交定位")
    add_markdown_body(doc, source)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def add_source_page(doc, page_number, lines, source_range):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"源代码登记页 {page_number:02d} / 30 · {source_range}")
    set_run_font(r, size=9, color=BLUE, bold=True)
    for idx, line in enumerate(lines, start=1):
        p = doc.add_paragraph(style="Code Block")
        r = p.add_run(f"{idx:03d} | {line}" if line else f"{idx:03d} | ")
        set_run_font(r, name="Consolas", size=6.5, color=INK)


def collect_source_lines():
    lines = []
    for group in ("front", "middle", "back"):
        folder = SOFTWARE / "source_code" / group
        for path in sorted(folder.glob("*.py")):
            lines.append(f"# ===== {group}/{path.name} =====")
            file_lines = path.read_text(encoding="utf-8").splitlines()
            lines.extend(file_lines)
            lines.append("")
    return lines


def build_source_printouts():
    all_lines = collect_source_lines()
    # The registration copy is a readable首/中/后 30-page excerpt. The
    # complete source remains in source_code/ and is indexed by the manifest.
    pages_per_doc = 30
    lines_per_page = 45
    excerpt_size = pages_per_doc * lines_per_page
    mid_start = max(0, (len(all_lines) - excerpt_size) // 2)
    samples = [
        ("前30页源代码.docx", all_lines[:excerpt_size], "完整源代码前段节选"),
        ("中30页源代码.docx", all_lines[mid_start:mid_start + excerpt_size], "完整源代码中段节选"),
        ("后30页源代码.docx", all_lines[-excerpt_size:], "完整源代码后段节选"),
    ]
    for name, excerpt, label in samples:
        doc = Document()
        configure_document(doc, "软件著作权 · 源代码提交件", preset="compact_reference_guide")
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        chunks = [excerpt[i * lines_per_page:(i + 1) * lines_per_page] for i in range(pages_per_doc)]
        for page_idx, chunk in enumerate(chunks):
            if page_idx > 0:
                doc.add_page_break()
            add_source_page(doc, page_idx + 1, chunk, label)
        path = WORD_SOURCE / name
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)


def write_svg(path, title, boxes, arrows, width=1200, height=700):
    box_svg = []
    for x, y, w, h, text, fill in boxes:
        lines = text.split("\n")
        box_svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="16" fill="#{fill}" stroke="#2E74B5" stroke-width="3"/>')
        for idx, line in enumerate(lines):
            box_svg.append(f'<text x="{x + w / 2}" y="{y + 55 + idx * 30}" text-anchor="middle" font-family="Microsoft YaHei, Arial" font-size="24" fill="#0B2545">{escape(line)}</text>')
    arrow_svg = []
    for x1, y1, x2, y2 in arrows:
        arrow_svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#2E74B5" stroke-width="4" marker-end="url(#arrow)"/>')
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">
<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#2E74B5"/></marker></defs>
<rect width="100%" height="100%" fill="#FFFFFF"/>
<text x="60" y="58" font-family="Microsoft YaHei, Arial" font-size="32" font-weight="700" fill="#0B2545">{escape(title)}</text>
{''.join(arrow_svg)}{''.join(box_svg)}
</svg>'''
    path.write_text(svg, encoding="utf-8")


def build_rendered_figures():
    FIGURES.mkdir(parents=True, exist_ok=True)
    write_svg(FIGURES / "fig01_system_architecture.svg", "双场景数字孪生与建议式决策架构", [
        (50, 180, 250, 150, "施工曲线\n压力 / 液量 / 砂比", LIGHT_BLUE),
        (475, 180, 250, 150, "PKN + EnKF\n在线状态估计", "E8F2FF"),
        (900, 180, 250, 150, "智能体建议\n安全门禁 / 人工确认", "E9F6F3"),
        (260, 470, 250, 130, "无 DAS\n压力校正", "F4F6F9"),
        (690, 470, 250, 130, "有 DAS\n六簇观测校验", "F4F6F9"),
    ], [(300, 255, 475, 255), (725, 255, 900, 255), (600, 330, 385, 470), (600, 330, 815, 470)] )
    write_svg(FIGURES / "fig02_kg_enkf_flow.svg", "KG-EnKF 先验桥接与参数更新", [
        (50, 190, 240, 140, "知识图谱规则\n工况 / 物性关系", LIGHT_BLUE),
        (365, 190, 240, 140, "先验桥接\n均值 + 协方差 + 置信度", "E8F2FF"),
        (680, 190, 240, 140, "PKN 正演\n得到压力与分簇响应", "E8F2FF"),
        (995, 190, 155, 140, "EnKF\n同化观测", "E9F6F3"),
        (520, 470, 300, 130, "后验参数\n下一时刻重新正演", "FFF8E8"),
    ], [(290, 260, 365, 260), (605, 260, 680, 260), (920, 260, 995, 260), (1070, 330, 670, 470)] )
    write_svg(FIGURES / "fig03_piggy_bank.svg", "阶段级 Piggy-Bank 液量再分配", [
        (50, 180, 260, 150, "阶段计划总液量\nV_stage", LIGHT_BLUE),
        (470, 180, 260, 150, "响应速率 / 均衡度\n识别快段与慢段", "E8F2FF"),
        (890, 180, 260, 150, "储备账本\nrelease / draw", "FFF8E8"),
        (260, 480, 260, 130, "下一窗降液量\n快段少分", "E9F6F3"),
        (680, 480, 260, 130, "下一窗增液量\n慢段多分", "E9F6F3"),
    ], [(310, 255, 470, 255), (730, 255, 890, 255), (1020, 330, 390, 480), (1020, 330, 810, 480)])
    write_svg(FIGURES / "fig04_safety_gate.svg", "建议式决策安全门禁", [
        (70, 210, 250, 140, "状态与观测\n压力 / 砂比 / 质量", LIGHT_BLUE),
        (475, 210, 250, 140, "策略候选\nTD3 / 规则 / 保持", "E8F2FF"),
        (880, 210, 250, 140, "硬约束检查\n风险优先", "FFF8E8"),
        (475, 480, 250, 140, "人工确认\n高风险需确认", "E9F6F3"),
        (880, 480, 250, 140, "输出建议\n不直接控制设备", "F4F6F9"),
    ], [(320, 280, 475, 280), (725, 280, 880, 280), (1005, 350, 600, 480), (1005, 350, 1005, 480)])


def build_checklist_doc():
    doc = Document()
    configure_document(doc, "申报包 · 递交检查表", preset="standard_business_brief")
    add_title_block(doc, "合同第二、三部分知识产权申报包", "正式递交前检查清单与材料边界", "申报材料递交检查表")
    add_callout(doc, "本清单将当前技术底稿、运行交付包和正式登记材料分开管理。标记为“待补”的项目需要由申请主体、代理师或登记人员补齐，不用技术代码虚构。")
    doc.add_heading("一、专利材料", level=1)
    rows = [
        ["材料", "当前状态", "递交动作"],
        ["技术交底书", "已生成 Word 正式整理版", "代理师结合现有技术检索后压缩或调整权利要求"],
        ["权利要求书初稿", "已形成技术草案", "由代理师完成法律表述、从属项取舍和新颖性审查"],
        ["实施例与实验依据", "已生成 Word 正式整理版", "核对每个指标的实际证据等级，避免把研究接口写成已验收"],
        ["说明书附图", "已提供 SVG 和 Mermaid 源图", "代理师按专利制图规范统一线宽、编号和图注"],
        ["申请主体与发明人", "待补", "填写单位、发明人、联系人和盖章信息"],
    ]
    add_table_from_markdown(doc, ["|" + "|".join(r) + "|" for r in rows])
    doc.add_heading("二、软件著作权材料", level=1)
    rows = [
        ["材料", "当前状态", "递交动作"],
        ["软件功能说明书", "已生成 Word 正式整理版", "核对软件名称、版本号和开发完成日期"],
        ["用户手册", "已生成 Word 正式整理版", "补充最终截图和甲方指定封面信息"],
        ["源代码提交件", "已生成前/中/后 30 页 Word 文件", "确认打印页数、连续性和代码脱敏范围"],
        ["申请表与审批单", "未代填", "使用单位正式模板，补齐主体信息后盖章"],
        ["运行包", "已保留独立 HMI 交付包", "运行包与登记材料分开归档，不将合成演示数据写成现场数据"],
    ]
    add_table_from_markdown(doc, ["|" + "|".join(r) + "|" for r in rows])
    doc.add_heading("三、交付包卫生检查", level=1)
    for item in [
        "删除 __pycache__、*.pyc、构建缓存和本机临时日志。",
        "确认原始施工数据、原始 DAS 数据、真实井轨迹和敏感配置不在公开包中。",
        "保留 source_code_manifest.csv、版本清单和脱敏说明，用于追溯提交边界。",
        "确认 Word 文件完成渲染检查后再交给代理师或登记人员。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, size=11, color=INK)
    out = DELIVERABLES / "00_申报材料目录与递交检查表.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def build_delivery_manifest():
    entries = []
    for path in sorted(DELIVERABLES.rglob("*")):
        if not path.is_file():
            continue
        if path.name == "申报包文件清单.json":
            continue
        if "__pycache__" in path.parts or path.suffix.lower() in {".pyc", ".pyo"}:
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        entries.append({
            "path": str(path.relative_to(ROOT)).replace("\\", "/"),
            "bytes": path.stat().st_size,
            "sha256": digest,
        })
    payload = {
        "generated_at": "2026-08-30",
        "scope": "patent_part2_part3 + software_copyright_part2_part3 + third_part_hmi_delivery",
        "excluded": ["__pycache__", "*.pyc", "*.pyo", "原始施工数据", "原始DAS数据", "真实现场配置"],
        "files": entries,
    }
    (DELIVERABLES / "申报包文件清单.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    WORD_PATENT.mkdir(parents=True, exist_ok=True)
    WORD_SOFTWARE.mkdir(parents=True, exist_ok=True)
    WORD_SOURCE.mkdir(parents=True, exist_ok=True)
    build_markdown_doc(
        PATENT / "01_专利技术交底书.md",
        WORD_PATENT / "01_专利技术交底书.docx",
        "知识引导参数同化与阶段级液量均衡调控方法",
        "合同第二、三部分 · 发明专利技术交底正式整理版",
        "专利技术交底书",
        intro="本文件是技术交底正式整理版，不替代代理师对现有技术、新颖性和权利要求保护范围的法律判断。",
    )
    build_markdown_doc(
        PATENT / "04_实施例与实验依据.md",
        WORD_PATENT / "04_实施例与实验依据.docx",
        "实施例与实验依据",
        "双场景数字孪生、KG-EnKF、Piggy-Bank 与建议式决策",
        "专利实施例",
        intro="实验指标按证据等级表述；对于仍处于研究或接口阶段的 PyFrac 长时动态、现场闭环和簇级直接控制，不写成已完成验收。",
    )
    build_markdown_doc(
        PATENT / "02_权利要求书初稿.md",
        WORD_PATENT / "02_权利要求书初稿.docx",
        "专利权利要求书初稿",
        "知识引导参数同化与阶段级总液量均衡调控",
        "权利要求书初稿",
        intro="本文件是技术草案，供代理师结合现有技术检索后修改，不作为最终法律文本。",
    )
    build_markdown_doc(
        PATENT / "03_说明书附图与绘图说明.md",
        WORD_PATENT / "03_说明书附图与绘图说明.docx",
        "说明书附图与绘图说明",
        "双场景数字孪生、KG-EnKF 与阶段级均衡调控",
        "说明书附图",
        intro="附图源文件位于 figures/，正式申请时由代理师按制图规范统一调整图号、线型和版式。",
    )
    build_markdown_doc(
        PATENT / "06_申请边界与代理师审查清单.md",
        WORD_PATENT / "06_申请边界与代理师审查清单.docx",
        "申请边界与代理师审查清单",
        "需要检索、补证和确认的事项",
        "代理师审查清单",
        intro="本清单用于把技术创新点、代码证据和仍待验证的内容分开，便于代理师组织保护范围。",
    )
    build_markdown_doc(
        PATENT / "08_专利补充材料_阶段级总液量均衡调控.md",
        WORD_PATENT / "08_专利补充材料_阶段级总液量均衡调控.docx",
        "专利补充材料：阶段级总液量均衡调控",
        "Piggy-Bank 账本、响应效率与开环验证路径",
        "专利补充材料",
        intro="本材料补充说明只能调节阶段总液量时的均衡调控路径，不把簇级估计量写成可直接执行的单簇液量控制。",
    )
    build_markdown_doc(
        PATENT / "07_申报信息模板.md",
        WORD_PATENT / "07_申报信息模板.docx",
        "专利申报信息模板",
        "申请主体、发明人和联系人信息待补表",
        "专利申报信息",
        intro="本页仅提供待填写字段，不代填申请人、发明人、日期和盖章信息。",
    )
    build_markdown_doc(
        SOFTWARE / "01_软件功能说明书.md",
        WORD_SOFTWARE / "01_软件功能说明书.docx",
        "智能压裂双场景数字孪生与安全建议软件 V1.0",
        "软件著作权功能说明正式整理版",
        "软件功能说明书",
        intro="软件定位为状态估计、解释性分析和建议式决策平台；PKN+EnKF 是默认在线模型，智能体建议经过安全门禁，不直接控制现场设备。",
    )
    build_markdown_doc(
        SOFTWARE / "02_用户手册.md",
        WORD_SOFTWARE / "02_用户手册.docx",
        "智能压裂双场景数字孪生与安全建议软件 V1.0",
        "软件著作权用户手册正式整理版",
        "用户手册",
        preset="compact_reference_guide",
        intro="本手册用于说明软件的启动、场景切换、结果查看和结果导出；申报前需要将最终脱敏界面截图补入截图位置。",
    )
    build_markdown_doc(
        SOFTWARE / "08_申报信息模板.md",
        WORD_SOFTWARE / "08_申报信息模板.docx",
        "软件著作权申报信息模板",
        "著作权人、版本和日期信息待补表",
        "软件著作权申报信息",
        intro="本页仅提供待填写字段，不代填著作权人、开发完成日期、首次发表日期和盖章信息。",
    )
    for filename, title, subtitle, label in [
        ("09_登记审批表填报模板.md", "软件著作权登记审批表填报模板", "主体信息、版本和审批字段待补", "登记审批表"),
        ("10_软件著作权申请表填报模板.md", "软件著作权申请表填报模板", "申请主体和软件摘要待补", "软件著作权申请表"),
        ("11_附件1_申报书填报模板.md", "软件著作权申报书填报模板", "功能模块和提交附件清单", "附件1申报书"),
        ("12_审批单填报模板.md", "软件著作权审批单填报模板", "内部审核、保密和盖章字段待补", "审批单"),
    ]:
        build_markdown_doc(
            SOFTWARE / filename,
            WORD_SOFTWARE / filename.replace(".md", ".docx"),
            title,
            subtitle,
            label,
            intro="本模板参考现有申报材料的组织方式，仅提供待填字段；正式递交以申请单位和登记机构的现行表单为准。",
        )
    build_source_printouts()
    build_rendered_figures()
    build_checklist_doc()
    build_delivery_manifest()
    print("Generated Word submission materials and rendered SVG figures.")
    for path in sorted(WORD_PATENT.glob("*.docx")):
        print(path)
    for path in sorted(WORD_SOFTWARE.glob("*.docx")):
        print(path)
    for path in sorted(WORD_SOURCE.glob("*.docx")):
        print(path)
    for path in sorted(FIGURES.glob("*.svg")):
        print(path)


if __name__ == "__main__":
    main()
