from pathlib import Path
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
# Keep the audit aligned with the IP package builder.  The repository root
# contains a separate, unrelated deliverables directory.
DELIVERABLES = ROOT / "专利与软著" / "deliverables"


def docx_audit(path):
    doc = Document(path)
    section = doc.sections[0]
    headings = {"Heading 1": 0, "Heading 2": 0, "Heading 3": 0}
    for paragraph in doc.paragraphs:
        if paragraph.style.name in headings:
            headings[paragraph.style.name] += 1
    page_breaks = 0
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        all_xml = "".join(
            data.decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
            for data in [archive.read(name)]
        )
        page_breaks = xml.count('w:br w:type="page"')
    return {
        "file": str(path.relative_to(ROOT)),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "headings": headings,
        "explicit_page_breaks": page_breaks,
        "page_width_in": round(section.page_width.inches, 2),
        "page_height_in": round(section.page_height.inches, 2),
        "margins_in": [round(section.top_margin.inches, 2), round(section.right_margin.inches, 2), round(section.bottom_margin.inches, 2), round(section.left_margin.inches, 2)],
        "has_page_field": "PAGE" in all_xml,
    }


def main():
    docs = [
        DELIVERABLES / "00_申报材料目录与递交检查表.docx",
        *(DELIVERABLES / "patent_part2_part3" / "word_submission").glob("*.docx"),
        *(DELIVERABLES / "software_copyright_part2_part3" / "word_submission").glob("*.docx"),
        *(DELIVERABLES / "software_copyright_part2_part3" / "source_code_submission").glob("*.docx"),
    ]
    print("DOCX_AUDIT")
    for path in sorted(docs):
        print(docx_audit(path))

    svgs = sorted((DELIVERABLES / "patent_part2_part3" / "figures" / "rendered").glob("*.svg"))
    print("SVG_AUDIT count=" + str(len(svgs)))
    for path in svgs:
        text = path.read_text(encoding="utf-8")
        print(path.name, "valid_svg=" + str(text.lstrip().startswith("<svg")), "bytes=" + str(path.stat().st_size))

    cache_files = [p for p in DELIVERABLES.rglob("*") if p.is_file() and (p.suffix.lower() in {".pyc", ".pyo"} or "__pycache__" in p.parts)]
    print("CACHE_FILES")
    for path in cache_files:
        print(path.relative_to(ROOT))
    print("CACHE_COUNT=" + str(len(cache_files)))


if __name__ == "__main__":
    main()
